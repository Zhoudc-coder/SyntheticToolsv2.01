import csv
import datetime
import json
import uuid
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook, Workbook

try:
    import xlsxwriter
    HAS_XLSXWRITER = True
except ImportError:
    HAS_XLSXWRITER = False

from config import (
    DB_PATH, RETENTION_DAYS, get_backup_dir, get_device_id,
    get_unify_version_when_ignore_box, get_summary_format_is_csv,
)
from db import LogDatabase
from file_parser import validate_sn


REASON_LABELS = {
    'no_version': '文件未提取到版本号',
    'date_mismatch': '文件日期与选定日期不匹配',
    'upload_time_mismatch': '文件上传时间不在指定范围',
    'box_code_duplicate_db': '箱码已存在于数据库',
    'box_code_duplicate_batch': '同批次箱码重复',
    'read_error': '文件读取失败',
    'no_date': '文件日期无法确定',
    'sn_format': 'SN格式不符合',
    'empty_sn': 'SN为空（其他列有数据）',
    'sn_duplicate_db': 'SN与历史记录重复',
    'sn_duplicate_batch': '同批次SN重复',
    'row_too_wide': '行数据列数超过标题（已截断）',
}

SKIP_STATS_FILE_CATS = [
    'no_version', 'date_mismatch',
    'box_code_duplicate_db', 'box_code_duplicate_batch',
    'read_error', 'no_date',
]

SKIP_STATS_SN_CATS = [
    'sn_format', 'empty_sn', 'sn_duplicate_db',
    'sn_duplicate_batch', 'row_too_wide',
]


def _read_rows(file_path: Path):
    suffix = file_path.suffix.lower()
    if suffix == '.xlsx':
        wb = load_workbook(filename=str(file_path), read_only=True, data_only=True)
        ws = wb.active
        try:
            title_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        except StopIteration:
            wb.close()
            raise ValueError("文件没有标题行")

        def data_iter():
            try:
                for row in ws.iter_rows(min_row=2, values_only=True):
                    yield row
            finally:
                wb.close()
        return title_row, data_iter()
    elif suffix == '.csv':
        encodings = ['utf-8-sig', 'utf-8', 'gbk', 'gb18030']
        for enc in encodings:
            try:
                f = open(file_path, 'r', encoding=enc, newline='')
                reader = csv.reader(f)
                try:
                    title_row = next(reader)
                except StopIteration:
                    f.close()
                    raise ValueError("CSV 文件为空")

                def data_iter():
                    try:
                        for row in reader:
                            yield row
                    finally:
                        f.close()
                return title_row, data_iter()
            except UnicodeDecodeError:
                continue
        raise ValueError("无法解码 CSV 文件，请使用 UTF-8 或 GBK 编码")
    else:
        raise ValueError(f"不支持的文件类型: {suffix}")


def _get_unique_filename(directory: Path, base_name: str, extension: str = '.xlsx') -> Path:
    final_path = directory / (base_name + extension)
    counter = 1
    while final_path.exists():
        final_path = directory / (f"{base_name} ({counter})" + extension)
        counter += 1
    return final_path


def _is_zero_value(v) -> bool:
    if v is None:
        return False
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return v == 0
    s = str(v).strip()
    if s == '':
        return False
    try:
        return float(s) == 0
    except (ValueError, TypeError):
        return False


def _is_abnormal_row(row) -> bool:
    if row is None or len(row) < 11:
        return False
    for i in range(1, 11):
        if not _is_zero_value(row[i]):
            return False
    return True


def _normalize_cell_value(v):
    import math
    if v is None:
        return ''
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return str(v)
        return v
    if isinstance(v, str):
        return v
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return v
    return str(v)


def _normalize_csv_cell_value(v):
    """CSV 专用：在 _normalize_cell_value 基础上，将布尔/日期等转为字符串，确保写入正确。"""
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'True' if v else 'False'
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return str(v)
    if isinstance(v, float):
        import math
        if math.isnan(v) or math.isinf(v):
            return str(v)
        return v
    return v


def _row_to_json(row) -> str:
    try:
        return json.dumps([str(x) if x is not None else '' for x in row], ensure_ascii=False)
    except Exception:
        return ''


def perform_merge(file_infos, output_dir: Path, output_filename_prefix: str,
                  log_func=print,
                  progress_callback=None,
                  auto_export_boxes: bool = True,
                  batch_id_override: str = None,
                  special_mode: bool = False,
                  upload_start: str = None,
                  upload_end: str = None,
                  ignore_sn_format: bool = False,
                  ignore_box_format: bool = False,
                  simple_mode: bool = False,
                  check_duplicates: bool = True,
                  scan_skipped_items=None,
                  upload_start_time: str = '',
                  upload_end_time: str = '',
                  force_merge: int = 0,
                  device_id: str = '') -> tuple[bool, list[str], str, str]:
    if not file_infos:
        return False, [], "没有待合并的文件", ""

    if not HAS_XLSXWRITER:
        log_func("警告：未安装 xlsxwriter，回退到 openpyxl 写入模式，处理大量宽表数据时可能内存不足。建议执行 pip install xlsxwriter。")

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    if batch_id_override:
        batch_id = batch_id_override
    else:
        batch_id = (datetime.datetime.now().strftime('%Y%m%d%H%M%S%f') + uuid.uuid4().hex)[:20]

    db = LogDatabase(DB_PATH)
    try:
        db.initialize()
        db.cleanup_old_logs(RETENTION_DAYS)
        existing_box_codes = set()
        existing_box_filename = {}
        for box_code, original_filename in db.conn.execute("SELECT box_code, original_filename FROM merged_boxes"):
            existing_box_codes.add(box_code)
            existing_box_filename[box_code] = original_filename
        existing_sns = set()
        sn_to_box_code = {}
        for sn, box_code in db.conn.execute("SELECT sn, box_code FROM merged_sns"):
            existing_sns.add(sn)
            sn_to_box_code[sn] = box_code
    except Exception as e:
        return False, [], f"日志数据库初始化失败：{e}", ""

    merge_date = datetime.date.today()
    merge_date_iso = merge_date.isoformat()

    # ===== 汇总文件格式配置 =====
    # summary_format_is_csv = False：与原版本一致，生成 xlsx 汇总文件
    # summary_format_is_csv = True ：生成 CSV 汇总文件
    summary_format_is_csv = get_summary_format_is_csv()
    if summary_format_is_csv:
        log_func("提示：根据配置 summary_format_is_csv=1，本次合并生成的汇总文件将使用 CSV 格式。")

    # ===== 分组逻辑 =====
    # 当启用了"忽略箱码格式"且配置项 unify_version_when_ignore_box 为 1 时：
    #   本次合并将忽略各文件的解析版本差异，统一归入一个分组，
    #   分组版本号取所有文件中出现次数最多的版本。
    # 其它情况保持原有分组逻辑。
    unify_version_when_ignore_box = get_unify_version_when_ignore_box()
    unified_mode = bool(ignore_box_format and unify_version_when_ignore_box)

    groups = {}
    if unified_mode:
        version_counter = Counter()
        date_counter = Counter()
        for _fp, _bi in file_infos:
            version_counter[_bi.version] += 1
            date_counter[_bi.package_date] += 1
        if not version_counter:
            return False, [], "没有待合并的文件", ""
        mode_version = version_counter.most_common(1)[0][0]
        mode_date = date_counter.most_common(1)[0][0]
        log_func("提示：已启用“忽略箱码格式”且配置项 unify_version_when_ignore_box=1")
        if special_mode:
            groups = {mode_version: list(file_infos)}
        else:
            groups = {(mode_date, mode_version): list(file_infos)}
    elif special_mode:
        for file_path, box_info in file_infos:
            groups.setdefault(box_info.version, []).append((file_path, box_info))
    else:
        for file_path, box_info in file_infos:
            groups.setdefault((box_info.package_date, box_info.version), []).append((file_path, box_info))

    total_files = len(file_infos)
    processed_files = 0
    total_groups = len(groups)
    group_index = 0

    box_records = []
    sn_records = []
    abnormal_records = []
    skipped_records = []
    seen_sns = set()
    seen_sn_source = {}
    generated_files = []
    group_stats = {}

    def _add_merge_skip(item_type, box_code, sn, source_file, reason_cat, reason_detail,
                        shipment_date, upload_time, row_data=''):
        ts = datetime.datetime.now()
        skipped_records.append((
            item_type, box_code, sn, source_file, reason_cat, reason_detail,
            shipment_date or '', upload_time or '',
            merge_date_iso, ts.strftime('%Y-%m-%d %H:%M'), row_data
        ))

    for group_key, files in groups.items():
        group_index += 1
        if special_mode:
            version = group_key
            date_key = None
            log_func(f"开始处理版本 v{version}（特殊模式），共 {len(files)} 个文件")
        else:
            date_key, version = group_key
            log_func(f"开始处理 日期 {date_key} 版本 v{version}，共 {len(files)} 个文件")
        group_stats[group_key] = {'files': 0, 'total_rows': 0, 'valid_rows': 0, 'abnormal_rows': 0}

        group_box_records = []

        if special_mode:
            start_short = upload_start.replace('-', '')[4:].replace(':', '').replace(' ', '') if upload_start else 'unknown'
            end_short = upload_end.replace('-', '')[4:].replace(':', '').replace(' ', '') if upload_end else 'unknown'
            base_filename = f"{output_filename_prefix}-{start_short}-{end_short}-v{version}"
        else:
            try:
                date_obj = datetime.date.fromisoformat(date_key)
                date_str_short = date_obj.strftime('%y%m%d')
            except:
                date_str_short = date_key.replace('-', '')[2:]
            base_filename = f"{output_filename_prefix}-{date_str_short}-v{version}"

        # ===== 根据配置决定输出扩展名 =====
        output_extension = '.csv' if summary_format_is_csv else '.xlsx'

        final_path = _get_unique_filename(output_dir, base_filename, output_extension)
        final_filename = final_path.name
        temp_path = output_dir / (final_filename + '.tmp')

        if simple_mode:
            title_row = ['SN', 'box_code', 'source_file']
            title_len = 3
        else:
            try:
                first_file = files[0][0]
                title_row, _ = _read_rows(first_file)
                title_len = len(title_row)
            except Exception as e:
                log_func(f"读取组 {base_filename} 第一个文件标题失败：{first_file.name} - {e}")
                continue

        # ===== 初始化写入器 =====
        csv_file = None
        csv_writer = None
        xw_workbook = None
        xw_worksheet = None
        xw_current_row = 0
        out_wb = None
        out_ws = None
        HAS_XLSXWRITER_LOCAL = False
        writer_init_failed = False

        if summary_format_is_csv:
            # CSV 分支
            try:
                csv_file = open(temp_path, 'w', encoding='utf-8-sig', newline='')
                csv_writer = csv.writer(csv_file, quoting=csv.QUOTE_MINIMAL)
                csv_writer.writerow([_normalize_csv_cell_value(v) for v in title_row])
            except Exception as e:
                log_func(f"初始化 CSV 写入器失败：{e}")
                if csv_file is not None:
                    try:
                        csv_file.close()
                    except Exception:
                        pass
                    csv_file = None
                    csv_writer = None
                writer_init_failed = True
            if writer_init_failed:
                try:
                    if temp_path.exists():
                        temp_path.unlink()
                except Exception:
                    pass
                continue
        else:
            # XLSX 分支（保持原逻辑）
            if HAS_XLSXWRITER:
                try:
                    xw_workbook = xlsxwriter.Workbook(str(temp_path), {'constant_memory': True, 'use_zip64': True})
                    xw_worksheet = xw_workbook.add_worksheet("汇总")
                    for col_idx, val in enumerate(title_row):
                        xw_worksheet.write(0, col_idx, _normalize_cell_value(val))
                    xw_current_row = 1
                    HAS_XLSXWRITER_LOCAL = True
                except Exception as e:
                    log_func(f"初始化 xlsxwriter 失败，回退到 openpyxl：{e}")
                    try:
                        if xw_workbook is not None:
                            xw_workbook.close()
                    except Exception:
                        pass
                    xw_workbook = None
                    xw_worksheet = None
                    HAS_XLSXWRITER_LOCAL = False

            if not HAS_XLSXWRITER_LOCAL:
                out_wb = Workbook(write_only=True)
                out_ws = out_wb.create_sheet(title="汇总")
                out_ws.append(list(title_row))

        for file_path, box_info in files:
            processed_files += 1
            log_func(f"处理文件：{file_path.name} (组 {base_filename})")
            if progress_callback:
                progress_callback(f"正在合并：{file_path.name}（{processed_files}/{total_files}）")

            data_iter = None
            # 每个文件内 "行数据列数超过标题" 只输出一次日志
            row_too_wide_logged = False
            # 每个文件内 "文件标题列数不一致" / "文件标题内容存在差异" 只输出一次日志
            file_title_mismatch_logged = False
            try:
                file_title, data_iter = _read_rows(file_path)

                if not simple_mode and not file_title_mismatch_logged:
                    if len(file_title) != title_len:
                        log_func(f"警告：文件标题列数不一致（{len(file_title)} 列 vs {title_len} 列），将按基准列数 {title_len} 对齐，继续处理：{file_path.name}")
                        file_title_mismatch_logged = True
                    elif any(a != b for a, b in zip(file_title, title_row)):
                        log_func(f"警告：文件标题内容存在差异，可能影响列对应关系，但将继续处理：{file_path.name}")
                        file_title_mismatch_logged = True

                file_total_rows = 0
                row_count = 0
                file_abnormal_count = 0
                merge_time_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M')

                try:
                    mtime = file_path.stat().st_mtime
                    upload_time = datetime.datetime.fromtimestamp(mtime).strftime('%Y-%m-%d %H:%M')
                except:
                    upload_time = ''

                shipment_date_for_record = box_info.package_date

                for row in data_iter:
                    if row is None:
                        continue
                    is_all_empty = True
                    for v in row:
                        if v is None:
                            continue
                        if isinstance(v, str):
                            if v.strip():
                                is_all_empty = False
                                break
                        else:
                            is_all_empty = False
                            break
                    if is_all_empty:
                        continue

                    file_total_rows += 1

                    sn_raw = row[0] if len(row) > 0 else None
                    sn = str(sn_raw).strip() if sn_raw is not None else ''

                    if not ignore_sn_format:
                        if not validate_sn(sn):
                            _add_merge_skip('sn', box_info.box_code, sn or '<空>', file_path.name,
                                            'sn_format', f'SN格式不符合：{sn or "<空>"}',
                                            shipment_date_for_record, upload_time, _row_to_json(row))
                            log_func(f"警告：SN格式不符合，跳过行：{sn or '<空>'}")
                            continue

                    if not sn:
                        _add_merge_skip('sn', box_info.box_code, '', file_path.name,
                                        'empty_sn', 'SN为空但其他列有数据',
                                        shipment_date_for_record, upload_time, _row_to_json(row))
                        log_func("警告：SN为空（其他列有数据），跳过行")
                        continue

                    if check_duplicates:
                        if sn in existing_sns or sn in seen_sns:
                            if sn in existing_sns:
                                hist_box = sn_to_box_code.get(sn, '未知箱码')
                                hist_file = existing_box_filename.get(hist_box, '未知文件')
                                _add_merge_skip('sn', box_info.box_code, sn, file_path.name,
                                                'sn_duplicate_db', f'历史箱码：{hist_box}，历史文件：{hist_file}',
                                                shipment_date_for_record, upload_time, _row_to_json(row))
                                log_func(f"警告：SN重复（与历史记录重复），跳过：{sn}，历史箱码：{hist_box}，历史文件：{hist_file}")
                            else:
                                src_file = seen_sn_source.get(sn, '未知文件')
                                _add_merge_skip('sn', box_info.box_code, sn, file_path.name,
                                                'sn_duplicate_batch', f'首次出现文件：{src_file}',
                                                shipment_date_for_record, upload_time, _row_to_json(row))
                                log_func(f"警告：SN重复（本次合并已存在），跳过：{sn}，首次出现文件：{src_file}")
                            continue

                    if _is_abnormal_row(row):
                        abnormal_records.append((
                            sn, box_info.box_code, file_path.name,
                            shipment_date_for_record, upload_time, merge_date_iso, merge_time_str,
                            _row_to_json(row)
                        ))
                        file_abnormal_count += 1
                        log_func(f"提示：检测到异常数据行（前10列均为0），SN：{sn}，来源：{file_path.name}")

                    if simple_mode:
                        out_row = [sn, box_info.box_code, file_path.name]
                    else:
                        row_list = list(row)
                        if len(row_list) < title_len:
                            row_list.extend([None] * (title_len - len(row_list)))
                        elif len(row_list) > title_len:
                            _add_merge_skip('sn', box_info.box_code, sn, file_path.name,
                                            'row_too_wide', f'原行有 {len(row_list)} 列，已截断到 {title_len} 列',
                                            shipment_date_for_record, upload_time, _row_to_json(row))
                            if not row_too_wide_logged:
                                log_func(f"警告：行数据列数超过标题，截断：{sn}")
                                row_too_wide_logged = True
                            row_list = row_list[:title_len]
                        out_row = row_list

                    # ===== 逐行写入（CSV / XLSX 分支） =====
                    if summary_format_is_csv:
                        try:
                            csv_writer.writerow([_normalize_csv_cell_value(v) for v in out_row])
                        except Exception as ex:
                            log_func(f"警告：写入 CSV 行失败，跳过：{sn}，错误：{ex}")
                            continue
                    elif HAS_XLSXWRITER_LOCAL:
                        cleaned = [_normalize_cell_value(v) for v in out_row]
                        try:
                            xw_worksheet.write_row(xw_current_row, 0, cleaned)
                        except Exception as ex:
                            log_func(f"警告：写入行失败，跳过：{sn}，错误：{ex}")
                            continue
                        xw_current_row += 1
                    else:
                        out_ws.append(out_row)

                    seen_sns.add(sn)
                    seen_sn_source[sn] = file_path.name
                    sn_records.append((sn, box_info.box_code, merge_date_iso, merge_time_str))
                    row_count += 1

                # 记录箱码时写入 extracted_version（该文件识别出的版本，独立于汇总文件的版本）
                record_tuple = (
                    box_info.box_code, file_path.name, shipment_date_for_record, upload_time,
                    merge_date_iso, merge_time_str, final_filename,
                    row_count, file_total_rows, box_info.sequence_no,
                    box_info.version
                )
                box_records.append(record_tuple)
                group_box_records.append(record_tuple)

                group_stats[group_key]['files'] += 1
                group_stats[group_key]['total_rows'] += file_total_rows
                group_stats[group_key]['valid_rows'] += row_count
                group_stats[group_key]['abnormal_rows'] += file_abnormal_count

                if file_abnormal_count > 0:
                    log_func(f"文件 {file_path.name} 共检测到 {file_abnormal_count} 条异常数据行")

            except Exception as e:
                log_func(f"错误：读取文件失败，跳过：{file_path.name} - {e}")
                ts = datetime.datetime.now()
                skipped_records.append((
                    'file', box_info.box_code if box_info else file_path.stem, '', file_path.name,
                    'read_error', f'{e}', '', '',
                    merge_date_iso, ts.strftime('%Y-%m-%d %H:%M'), ''
                ))
                continue
            finally:
                if data_iter is not None:
                    try:
                        data_iter.close()
                    except Exception:
                        pass

        # ===== 关闭写入器：CSV 必须先 flush + close，保证数据完整落盘 =====
        save_ok = False
        try:
            if summary_format_is_csv:
                if csv_file is not None:
                    csv_file.flush()
                    csv_file.close()
                    csv_file = None
                save_ok = True
            elif HAS_XLSXWRITER_LOCAL:
                xw_workbook.close()
                save_ok = True
            else:
                out_wb.save(str(temp_path))
                save_ok = True
            log_func(f"组 {base_filename} 临时汇总文件已生成：{temp_path}")
        except Exception as e:
            log_func(f"组 {base_filename} 写入临时汇总文件失败：{e}")
            if summary_format_is_csv:
                if csv_file is not None:
                    try:
                        csv_file.close()
                    except Exception:
                        pass
                    csv_file = None
            else:
                try:
                    if HAS_XLSXWRITER_LOCAL and xw_workbook is not None:
                        xw_workbook.close()
                except Exception:
                    pass
            try:
                if temp_path.exists():
                    temp_path.unlink()
            except Exception:
                pass

        if not save_ok:
            continue

        try:
            if temp_path.exists():
                if final_path.exists():
                    final_path.unlink()
                temp_path.rename(final_path)
            generated_files.append(final_filename)
            log_func(f"组 {base_filename} 汇总完成：{final_filename}，有效 {group_stats[group_key]['valid_rows']} 行，总 {group_stats[group_key]['total_rows']} 行")
        except Exception as e:
            log_func(f"组 {base_filename} 替换最终文件失败：{e}")
            try:
                if temp_path.exists():
                    temp_path.unlink()
            except Exception:
                pass
            continue

        if auto_export_boxes and group_box_records:
            if progress_callback:
                progress_callback(f"正在自动导出箱码记录：{group_index}/{total_groups}")
            box_record_base = f"{base_filename}-箱码记录"
            box_record_path = _get_unique_filename(output_dir, box_record_base, '.xlsx')
            try:
                wb_box = Workbook(write_only=True)
                ws_box = wb_box.create_sheet("箱码记录")
                ws_box.append(["box_code", "original_filename", "shipment_date", "upload_time",
                               "merge_date", "merge_time", "output_filename",
                               "valid_rows", "total_rows", "sequence_no", "extracted_version"])
                for rec in group_box_records:
                    ws_box.append(list(rec))
                wb_box.save(str(box_record_path))
                log_func(f"箱码记录文件已生成：{box_record_path.name}")
                generated_files.append(box_record_path.name)
            except Exception as e:
                log_func(f"组 {base_filename} 生成箱码记录失败：{e}")

    if not generated_files:
        return False, [], "没有生成任何汇总文件（可能所有组均处理失败）", ""

    all_skipped = []
    if scan_skipped_items:
        all_skipped.extend(scan_skipped_items)
    all_skipped.extend(skipped_records)

    all_skip_stats = {}
    for rec in all_skipped:
        cat = rec[4]
        if cat == 'upload_time_mismatch':
            continue
        all_skip_stats[cat] = all_skip_stats.get(cat, 0) + 1

    if special_mode:
        ins_upload_start = upload_start_time or ''
        ins_upload_end = upload_end_time or ''
        ins_force_merge = int(force_merge)
    else:
        ins_upload_start = ''
        ins_upload_end = ''
        ins_force_merge = 0
    ins_device_id = device_id or get_device_id()

    try:
        insert_ok = db.insert_merge_records(
            box_records, sn_records, batch_id,
            abnormal_records, all_skipped,
            upload_start_time=ins_upload_start,
            upload_end_time=ins_upload_end,
            force_merge=ins_force_merge,
            device_id=ins_device_id
        )
    except Exception as e:
        for fname in generated_files:
            (output_dir / fname).unlink(missing_ok=True)
        return False, [], f"日志写入异常：{e}", ""

    if not insert_ok:
        for fname in generated_files:
            (output_dir / fname).unlink(missing_ok=True)
        return False, [], "日志写入失败：箱码或SN与已有记录冲突（可能其他电脑已合并），请重新操作", ""

    # ===== 写入共享盘备份数据库 =====
    backup_dir_str = get_backup_dir()
    if backup_dir_str:
        try:
            backup_root = Path(backup_dir_str)
            backup_path = backup_root / ins_device_id / 'merge_log.db'
            backup_path.parent.mkdir(parents=True, exist_ok=True)
            backup_db = LogDatabase(str(backup_path), is_backup=True)
            backup_db.initialize()
            backup_ok = backup_db.insert_merge_records(
                box_records, sn_records, batch_id,
                abnormal_records, all_skipped,
                upload_start_time=ins_upload_start,
                upload_end_time=ins_upload_end,
                force_merge=ins_force_merge,
                device_id=ins_device_id
            )
            backup_db.close()
            if backup_ok:
                log_func(f"备份数据库写入成功：{backup_path}")
            else:
                log_func(f"警告：备份数据库写入失败（可能记录已存在）：{backup_path}")
        except Exception as e:
            log_func(f"警告：写入共享盘备份数据库失败：{e}")

    total_scanned_files = len(file_infos)
    total_valid_files = sum(stats['files'] for stats in group_stats.values())
    total_rows_all = sum(stats['total_rows'] for stats in group_stats.values())
    total_valid_rows = sum(stats['valid_rows'] for stats in group_stats.values())
    total_abnormal_rows = sum(stats['abnormal_rows'] for stats in group_stats.values())

    summary_lines = [
        f"合并成功！共扫描 {total_scanned_files} 个文件，其中有效文件 {total_valid_files} 个。总共 {total_rows_all} 行数据，其中有效数据 {total_valid_rows} 行。生成 {len(generated_files)} 个文件。"
    ]
    if total_abnormal_rows > 0:
        summary_lines.append(f"检测到异常数据行 {total_abnormal_rows} 条（SN正常但前10列均为0），已单独记录。")

    if special_mode:
        for version, stats in sorted(group_stats.items()):
            summary_lines.append(f"  版本 v{version}: 文件数 {stats['files']}，总共 {stats['total_rows']} 行，有效数据 {stats['valid_rows']} 行，异常 {stats['abnormal_rows']} 行")
    else:
        for (date_key, version), stats in sorted(group_stats.items()):
            summary_lines.append(f"  日期 {date_key} 版本 v{version}: 文件数 {stats['files']}，总共 {stats['total_rows']} 行，有效数据 {stats['valid_rows']} 行，异常 {stats['abnormal_rows']} 行")

    summary_msg = "\n".join(summary_lines)

    skip_stats_msg = ""
    if all_skip_stats:
        skip_lines = ["跳过数据统计："]
        for cat in SKIP_STATS_FILE_CATS:
            if cat in all_skip_stats:
                skip_lines.append(f"  因{REASON_LABELS.get(cat, cat)}：跳过 {all_skip_stats[cat]} 个文件")
        for cat in SKIP_STATS_SN_CATS:
            if cat in all_skip_stats:
                skip_lines.append(f"  因{REASON_LABELS.get(cat, cat)}：跳过 {all_skip_stats[cat]} 行")
        if len(skip_lines) > 1:
            skip_stats_msg = "\n".join(skip_lines)

    log_func(summary_msg)
    if skip_stats_msg:
        log_func("")
        log_func(skip_stats_msg)

    finish_time_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    log_func("")
    log_func("=" * 60)
    log_func(f"合并完成时间：{finish_time_str}")
    log_func("=" * 60)

    return True, generated_files, summary_msg, skip_stats_msg