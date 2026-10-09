import sqlite3
import datetime
import os
import csv
import json
import re
from pathlib import Path
from openpyxl import load_workbook, Workbook
from config import DB_PATH, RETENTION_DAYS, ensure_log_dir


class LogDatabaseError(Exception):
    pass


class LogDatabase:
    def __init__(self, db_path=DB_PATH, is_backup: bool = False):
        self.db_path = str(db_path)
        self.is_backup = is_backup
        self.conn = None

    def connect(self):
        try:
            ensure_log_dir()
        except RuntimeError as e:
            raise LogDatabaseError(str(e))
        try:
            self.conn = sqlite3.connect(self.db_path, timeout=10)
            self.conn.execute("PRAGMA busy_timeout = 10000")
            try:
                self.conn.execute("PRAGMA journal_mode=WAL")
            except sqlite3.OperationalError:
                self.conn.execute("PRAGMA journal_mode=DELETE")
            self.conn.execute("PRAGMA synchronous=NORMAL")
            self.conn.execute("PRAGMA foreign_keys=ON")
        except sqlite3.Error as e:
            raise LogDatabaseError(f"无法连接日志数据库 {self.db_path}：{e}")

    def close(self):
        if self.conn:
            self.conn.close()
            self.conn = None

    def initialize(self):
        if not self.conn:
            self.connect()
        try:
            self.conn.executescript("""
                CREATE TABLE IF NOT EXISTS merged_boxes (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    batch_id TEXT NOT NULL DEFAULT '',
                    box_code TEXT NOT NULL,
                    original_filename TEXT NOT NULL,
                    shipment_date TEXT NOT NULL,
                    upload_time TEXT NOT NULL DEFAULT '',
                    merge_date TEXT NOT NULL,
                    merge_time TEXT NOT NULL,
                    output_filename TEXT NOT NULL,
                    valid_rows INTEGER NOT NULL,
                    total_rows INTEGER NOT NULL,
                    sequence_no INTEGER,
                    extracted_version TEXT NOT NULL DEFAULT ''
                );
                CREATE TABLE IF NOT EXISTS merged_sns (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    batch_id TEXT NOT NULL DEFAULT '',
                    sn TEXT NOT NULL,
                    box_code TEXT NOT NULL,
                    merge_date TEXT NOT NULL,
                    merge_time TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS merge_batches (
                    batch_id TEXT PRIMARY KEY,
                    merge_timestamp TEXT NOT NULL,
                    box_codes TEXT NOT NULL,
                    box_count INTEGER NOT NULL DEFAULT 0,
                    valid_sn_count INTEGER NOT NULL DEFAULT 0,
                    total_sn_count INTEGER NOT NULL DEFAULT 0,
                    upload_start_time TEXT NOT NULL DEFAULT '',
                    upload_end_time TEXT NOT NULL DEFAULT '',
                    force_merge INTEGER NOT NULL DEFAULT 0,
                    device_id TEXT NOT NULL DEFAULT ''
                );
                CREATE TABLE IF NOT EXISTS abnormal_data (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    batch_id TEXT NOT NULL DEFAULT '',
                    sn TEXT NOT NULL,
                    box_code TEXT NOT NULL,
                    original_filename TEXT NOT NULL,
                    shipment_date TEXT NOT NULL,
                    upload_time TEXT NOT NULL DEFAULT '',
                    merge_date TEXT NOT NULL,
                    merge_time TEXT NOT NULL,
                    row_data TEXT
                );
                CREATE TABLE IF NOT EXISTS skipped_items (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    batch_id TEXT NOT NULL DEFAULT '',
                    item_type TEXT NOT NULL,
                    box_code TEXT NOT NULL DEFAULT '',
                    sn TEXT NOT NULL DEFAULT '',
                    source_file TEXT NOT NULL DEFAULT '',
                    reason_category TEXT NOT NULL,
                    reason_detail TEXT NOT NULL DEFAULT '',
                    shipment_date TEXT NOT NULL DEFAULT '',
                    upload_time TEXT NOT NULL DEFAULT '',
                    merge_date TEXT NOT NULL DEFAULT '',
                    merge_time TEXT NOT NULL DEFAULT '',
                    row_data TEXT
                );
            """)
            self.conn.commit()

            if self.check_schema_status():
                ok, msg = self.migrate_schema()
                if not ok:
                    raise LogDatabaseError(f"数据库结构迁移失败：{msg}")

            self.conn.executescript("""
                CREATE INDEX IF NOT EXISTS idx_merged_boxes_batch_id ON merged_boxes(batch_id);
                CREATE INDEX IF NOT EXISTS idx_merged_boxes_box_code ON merged_boxes(box_code);
                CREATE INDEX IF NOT EXISTS idx_merged_boxes_merge_date ON merged_boxes(merge_date);
                CREATE INDEX IF NOT EXISTS idx_merged_boxes_shipment_date ON merged_boxes(shipment_date);
                CREATE INDEX IF NOT EXISTS idx_merged_boxes_upload_time ON merged_boxes(upload_time);
                CREATE INDEX IF NOT EXISTS idx_merged_sns_batch_id ON merged_sns(batch_id);
                CREATE INDEX IF NOT EXISTS idx_merged_sns_sn ON merged_sns(sn);
                CREATE INDEX IF NOT EXISTS idx_merged_sns_box_code ON merged_sns(box_code);
                CREATE INDEX IF NOT EXISTS idx_merged_sns_merge_date ON merged_sns(merge_date);
                CREATE INDEX IF NOT EXISTS idx_abnormal_data_batch_id ON abnormal_data(batch_id);
                CREATE INDEX IF NOT EXISTS idx_abnormal_data_sn ON abnormal_data(sn);
                CREATE INDEX IF NOT EXISTS idx_abnormal_data_shipment_date ON abnormal_data(shipment_date);
                CREATE INDEX IF NOT EXISTS idx_abnormal_data_upload_time ON abnormal_data(upload_time);
                CREATE INDEX IF NOT EXISTS idx_abnormal_data_merge_date ON abnormal_data(merge_date);
                CREATE INDEX IF NOT EXISTS idx_skipped_items_batch_id ON skipped_items(batch_id);
                CREATE INDEX IF NOT EXISTS idx_skipped_items_reason_category ON skipped_items(reason_category);
                CREATE INDEX IF NOT EXISTS idx_skipped_items_shipment_date ON skipped_items(shipment_date);
                CREATE INDEX IF NOT EXISTS idx_skipped_items_upload_time ON skipped_items(upload_time);
                CREATE INDEX IF NOT EXISTS idx_skipped_items_merge_date ON skipped_items(merge_date);
            """)
            self.conn.commit()
        except sqlite3.Error as e:
            raise LogDatabaseError(f"初始化数据库表失败：{e}")
        except LogDatabaseError:
            raise
        except Exception as e:
            raise LogDatabaseError(f"初始化数据库表失败：{e}")

    def check_schema_status(self) -> bool:
        if not self.conn:
            self.connect()
        cursor = self.conn.execute("PRAGMA table_info(merged_boxes)")
        columns = {row[1] for row in cursor.fetchall()}
        if 'id' not in columns:
            return True
        if 'batch_id' not in columns:
            return True
        if 'upload_time' not in columns:
            return True
        if 'extracted_version' not in columns:
            return True
        if 'valid_rows' not in columns and 'row_count' in columns:
            return True
        cursor = self.conn.execute("PRAGMA table_info(merged_sns)")
        sns_columns = {row[1] for row in cursor.fetchall()}
        if 'id' not in sns_columns or 'batch_id' not in sns_columns or 'merge_time' not in sns_columns:
            return True
        cursor = self.conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='merge_batches'")
        if cursor.fetchone():
            cursor = self.conn.execute("PRAGMA table_info(merge_batches)")
            batch_columns = {row[1] for row in cursor.fetchall()}
            required = {
                'box_count', 'valid_sn_count', 'total_sn_count',
                'upload_start_time', 'upload_end_time', 'force_merge', 'device_id'
            }
            if not required.issubset(batch_columns):
                return True
        else:
            return True
        cursor = self.conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='abnormal_data'")
        if not cursor.fetchone():
            return True
        cursor = self.conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='skipped_items'")
        if not cursor.fetchone():
            return True
        return False

    def migrate_schema(self) -> tuple[bool, str]:
        if not self.conn:
            self.connect()
        try:
            if not self.check_schema_status():
                return True, "数据库已是最新结构，无需更新"
            self.conn.execute("PRAGMA foreign_keys=OFF")

            migration_batch_id = 'MIG_' + datetime.datetime.now().strftime('%Y%m%d%H%M%S')

            cursor = self.conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='merged_boxes'")
            if cursor.fetchone():
                cursor = self.conn.execute("SELECT * FROM merged_boxes")
                col_names = [d[0] for d in cursor.description]
                old_rows = [dict(zip(col_names, row)) for row in cursor.fetchall()]
                self.conn.execute("DROP TABLE merged_boxes")
            else:
                old_rows = []

            self.conn.execute("""
                CREATE TABLE merged_boxes (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    batch_id TEXT NOT NULL DEFAULT '',
                    box_code TEXT NOT NULL,
                    original_filename TEXT NOT NULL,
                    shipment_date TEXT NOT NULL,
                    upload_time TEXT NOT NULL DEFAULT '',
                    merge_date TEXT NOT NULL,
                    merge_time TEXT NOT NULL,
                    output_filename TEXT NOT NULL,
                    valid_rows INTEGER NOT NULL,
                    total_rows INTEGER NOT NULL,
                    sequence_no INTEGER,
                    extracted_version TEXT NOT NULL DEFAULT ''
                )
            """)
            for r in old_rows:
                self.conn.execute(
                    "INSERT INTO merged_boxes "
                    "(batch_id, box_code, original_filename, shipment_date, upload_time, merge_date, merge_time, "
                    " output_filename, valid_rows, total_rows, sequence_no, extracted_version) "
                    "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                    (migration_batch_id, r.get('box_code', ''), r.get('original_filename', ''),
                     r.get('shipment_date', r.get('package_date', '')), r.get('upload_time', ''),
                     r.get('merge_date', ''), r.get('merge_time', ''), r.get('output_filename', ''),
                     r.get('valid_rows', r.get('row_count', 0)),
                     r.get('total_rows', r.get('row_count', 0)), r.get('sequence_no', 0),
                     r.get('extracted_version', ''))
                )

            cursor = self.conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='merged_sns'")
            if cursor.fetchone():
                cursor = self.conn.execute("SELECT * FROM merged_sns")
                col_names = [d[0] for d in cursor.description]
                sns_old_rows = [dict(zip(col_names, row)) for row in cursor.fetchall()]
                self.conn.execute("DROP TABLE merged_sns")
            else:
                sns_old_rows = []

            self.conn.execute("""
                CREATE TABLE merged_sns (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    batch_id TEXT NOT NULL DEFAULT '',
                    sn TEXT NOT NULL,
                    box_code TEXT NOT NULL,
                    merge_date TEXT NOT NULL,
                    merge_time TEXT NOT NULL
                )
            """)
            for r in sns_old_rows:
                self.conn.execute(
                    "INSERT INTO merged_sns (batch_id, sn, box_code, merge_date, merge_time) VALUES (?,?,?,?,?)",
                    (migration_batch_id, r.get('sn', ''), r.get('box_code', ''),
                     r.get('merge_date', ''), r.get('merge_time', ''))
                )

            cursor = self.conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='merge_batches'")
            if cursor.fetchone():
                cursor = self.conn.execute("SELECT * FROM merge_batches")
                col_names = [d[0] for d in cursor.description]
                batch_old_rows = [dict(zip(col_names, row)) for row in cursor.fetchall()]
                self.conn.execute("DROP TABLE merge_batches")
            else:
                batch_old_rows = []

            self.conn.execute("""
                CREATE TABLE merge_batches (
                    batch_id TEXT PRIMARY KEY,
                    merge_timestamp TEXT NOT NULL,
                    box_codes TEXT NOT NULL,
                    box_count INTEGER NOT NULL DEFAULT 0,
                    valid_sn_count INTEGER NOT NULL DEFAULT 0,
                    total_sn_count INTEGER NOT NULL DEFAULT 0,
                    upload_start_time TEXT NOT NULL DEFAULT '',
                    upload_end_time TEXT NOT NULL DEFAULT '',
                    force_merge INTEGER NOT NULL DEFAULT 0,
                    device_id TEXT NOT NULL DEFAULT ''
                )
            """)
            for r in batch_old_rows:
                self.conn.execute(
                    "INSERT INTO merge_batches "
                    "(batch_id, merge_timestamp, box_codes, box_count, valid_sn_count, total_sn_count, "
                    " upload_start_time, upload_end_time, force_merge, device_id) "
                    "VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (r.get('batch_id', ''), r.get('merge_timestamp', ''), r.get('box_codes', '[]'),
                     r.get('box_count', 0), r.get('valid_sn_count', 0), r.get('total_sn_count', 0),
                     r.get('upload_start_time', ''), r.get('upload_end_time', ''),
                     r.get('force_merge', 0), r.get('device_id', ''))
                )

            self.conn.execute("""
                CREATE TABLE IF NOT EXISTS abnormal_data (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    batch_id TEXT NOT NULL DEFAULT '',
                    sn TEXT NOT NULL,
                    box_code TEXT NOT NULL,
                    original_filename TEXT NOT NULL,
                    shipment_date TEXT NOT NULL,
                    upload_time TEXT NOT NULL DEFAULT '',
                    merge_date TEXT NOT NULL,
                    merge_time TEXT NOT NULL,
                    row_data TEXT
                )
            """)

            self.conn.execute("""
                CREATE TABLE IF NOT EXISTS skipped_items (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    batch_id TEXT NOT NULL DEFAULT '',
                    item_type TEXT NOT NULL,
                    box_code TEXT NOT NULL DEFAULT '',
                    sn TEXT NOT NULL DEFAULT '',
                    source_file TEXT NOT NULL DEFAULT '',
                    reason_category TEXT NOT NULL,
                    reason_detail TEXT NOT NULL DEFAULT '',
                    shipment_date TEXT NOT NULL DEFAULT '',
                    upload_time TEXT NOT NULL DEFAULT '',
                    merge_date TEXT NOT NULL DEFAULT '',
                    merge_time TEXT NOT NULL DEFAULT '',
                    row_data TEXT
                )
            """)

            self.conn.execute("PRAGMA foreign_keys=ON")
            self.conn.commit()
            return True, "数据库结构更新成功"
        except Exception as e:
            self.conn.rollback()
            try:
                self.conn.execute("PRAGMA foreign_keys=ON")
            except:
                pass
            return False, f"数据库迁移失败：{e}"

    def cleanup_old_logs(self, days=RETENTION_DAYS):
        if self.is_backup:
            return
        if not self.conn:
            self.connect()
        cutoff = (datetime.date.today() - datetime.timedelta(days=days)).isoformat()
        self.conn.execute("DELETE FROM merged_sns WHERE merge_date < ?", (cutoff,))
        self.conn.execute("DELETE FROM merged_boxes WHERE merge_date < ?", (cutoff,))
        self.conn.execute("DELETE FROM abnormal_data WHERE merge_date < ?", (cutoff,))
        self.conn.execute("DELETE FROM skipped_items WHERE merge_date < ?", (cutoff,))
        self.conn.commit()

    def load_existing_box_codes(self):
        if not self.conn:
            self.connect()
        cursor = self.conn.execute("SELECT DISTINCT box_code FROM merged_boxes")
        return {row[0] for row in cursor.fetchall()}

    def load_existing_sns(self):
        if not self.conn:
            self.connect()
        cursor = self.conn.execute("SELECT DISTINCT sn FROM merged_sns")
        return {row[0] for row in cursor.fetchall()}

    def insert_merge_records(self, box_records, sn_records, batch_id: str,
                             abnormal_records=None, skipped_records=None,
                             upload_start_time: str = '', upload_end_time: str = '',
                             force_merge: int = 0, device_id: str = ''):
        if not self.conn:
            self.connect()
        if abnormal_records is None:
            abnormal_records = []
        if skipped_records is None:
            skipped_records = []
        try:
            self.conn.execute("BEGIN IMMEDIATE")
            self.conn.executemany(
                "INSERT INTO merged_boxes "
                "(batch_id, box_code, original_filename, shipment_date, upload_time, merge_date, merge_time, "
                " output_filename, valid_rows, total_rows, sequence_no, extracted_version) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                [(batch_id,) + tuple(rec) for rec in box_records]
            )
            self.conn.executemany(
                "INSERT INTO merged_sns (batch_id, sn, box_code, merge_date, merge_time) VALUES (?,?,?,?,?)",
                [(batch_id,) + tuple(rec) for rec in sn_records]
            )
            if abnormal_records:
                self.conn.executemany(
                    "INSERT INTO abnormal_data "
                    "(batch_id, sn, box_code, original_filename, shipment_date, upload_time, merge_date, merge_time, row_data) "
                    "VALUES (?,?,?,?,?,?,?,?,?)",
                    [(batch_id,) + tuple(rec) for rec in abnormal_records]
                )
            if skipped_records:
                self.conn.executemany(
                    "INSERT INTO skipped_items "
                    "(batch_id, item_type, box_code, sn, source_file, reason_category, reason_detail, "
                    " shipment_date, upload_time, merge_date, merge_time, row_data) "
                    "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                    [(batch_id,) + tuple(rec) for rec in skipped_records]
                )
            box_codes_list = [rec[0] for rec in box_records]
            box_count = len(box_codes_list)
            valid_sn_count = len(sn_records)
            total_sn_count = sum(rec[8] for rec in box_records)
            merge_timestamp = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            self.conn.execute(
                "INSERT INTO merge_batches "
                "(batch_id, merge_timestamp, box_codes, box_count, valid_sn_count, total_sn_count, "
                " upload_start_time, upload_end_time, force_merge, device_id) "
                "VALUES (?,?,?,?,?,?,?,?,?,?)",
                (batch_id, merge_timestamp, json.dumps(box_codes_list),
                 box_count, valid_sn_count, total_sn_count,
                 upload_start_time or '', upload_end_time or '',
                 int(force_merge), device_id or '')
            )
            self.conn.commit()
            return True
        except sqlite3.IntegrityError:
            self.conn.rollback()
            return False
        except Exception:
            self.conn.rollback()
            raise

    def get_last_upload_end_time(self, device_id: str = None) -> str | None:
        if not self.conn:
            self.connect()
        if device_id:
            cursor = self.conn.execute(
                "SELECT upload_end_time FROM merge_batches "
                "WHERE upload_end_time IS NOT NULL AND upload_end_time != '' AND device_id = ? "
                "ORDER BY merge_timestamp DESC LIMIT 1",
                (device_id,)
            )
        else:
            cursor = self.conn.execute(
                "SELECT upload_end_time FROM merge_batches "
                "WHERE upload_end_time IS NOT NULL AND upload_end_time != '' "
                "ORDER BY merge_timestamp DESC LIMIT 1"
            )
        row = cursor.fetchone()
        return row[0] if row else None

    def delete_merge_records(self, box_codes):
        if not self.conn:
            self.connect()
        try:
            self.conn.execute("BEGIN IMMEDIATE")
            self.conn.executemany(
                "DELETE FROM merged_boxes WHERE box_code = ?",
                [(bc,) for bc in box_codes]
            )
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise

    def get_latest_batch(self):
        if not self.conn:
            self.connect()
        cursor = self.conn.execute(
            "SELECT batch_id, box_codes FROM merge_batches ORDER BY merge_timestamp DESC LIMIT 1"
        )
        row = cursor.fetchone()
        if not row:
            return None
        return {'batch_id': row[0], 'box_codes': json.loads(row[1])}

    def undo_last_merge(self) -> int:
        if not self.conn:
            self.connect()
        cursor = self.conn.execute(
            "SELECT batch_id FROM merge_batches ORDER BY merge_timestamp DESC LIMIT 1"
        )
        row = cursor.fetchone()
        if not row:
            return 0
        batch_id = row[0]
        try:
            self.conn.execute("BEGIN IMMEDIATE")
            cursor = self.conn.execute("DELETE FROM merged_boxes WHERE batch_id = ?", (batch_id,))
            deleted_count = cursor.rowcount
            self.conn.execute("DELETE FROM merged_sns WHERE batch_id = ?", (batch_id,))
            self.conn.execute("DELETE FROM abnormal_data WHERE batch_id = ?", (batch_id,))
            self.conn.execute("DELETE FROM skipped_items WHERE batch_id = ?", (batch_id,))
            self.conn.execute("DELETE FROM merge_batches WHERE batch_id = ?", (batch_id,))
            self.conn.commit()
            return deleted_count
        except Exception:
            self.conn.rollback()
            raise

    def undo_batch_by_id(self, batch_id: str) -> int:
        if not self.conn:
            self.connect()
        cursor = self.conn.execute(
            "SELECT batch_id FROM merge_batches WHERE batch_id = ?", (batch_id,)
        )
        if not cursor.fetchone():
            return 0
        try:
            self.conn.execute("BEGIN IMMEDIATE")
            cursor = self.conn.execute("DELETE FROM merged_boxes WHERE batch_id = ?", (batch_id,))
            deleted_count = cursor.rowcount
            self.conn.execute("DELETE FROM merged_sns WHERE batch_id = ?", (batch_id,))
            self.conn.execute("DELETE FROM abnormal_data WHERE batch_id = ?", (batch_id,))
            self.conn.execute("DELETE FROM skipped_items WHERE batch_id = ?", (batch_id,))
            self.conn.execute("DELETE FROM merge_batches WHERE batch_id = ?", (batch_id,))
            self.conn.commit()
            return deleted_count
        except Exception:
            self.conn.rollback()
            raise

    def count_batches_in_range(self, start_time: str, end_time: str) -> int:
        if not self.conn:
            self.connect()
        cursor = self.conn.execute(
            "SELECT COUNT(*) FROM merge_batches WHERE merge_timestamp >= ? AND merge_timestamp <= ?",
            (start_time, end_time)
        )
        return cursor.fetchone()[0]

    def undo_batches_in_range(self, start_time: str, end_time: str) -> int:
        if not self.conn:
            self.connect()
        cursor = self.conn.execute(
            "SELECT batch_id FROM merge_batches WHERE merge_timestamp >= ? AND merge_timestamp <= ?",
            (start_time, end_time)
        )
        batch_ids = [row[0] for row in cursor.fetchall()]
        if not batch_ids:
            return 0
        try:
            self.conn.execute("BEGIN IMMEDIATE")
            placeholders = ','.join(['?'] * len(batch_ids))
            self.conn.execute(f"DELETE FROM merged_boxes WHERE batch_id IN ({placeholders})", batch_ids)
            self.conn.execute(f"DELETE FROM merged_sns WHERE batch_id IN ({placeholders})", batch_ids)
            self.conn.execute(f"DELETE FROM abnormal_data WHERE batch_id IN ({placeholders})", batch_ids)
            self.conn.execute(f"DELETE FROM skipped_items WHERE batch_id IN ({placeholders})", batch_ids)
            self.conn.execute(f"DELETE FROM merge_batches WHERE batch_id IN ({placeholders})", batch_ids)
            self.conn.commit()
            return len(batch_ids)
        except Exception:
            self.conn.rollback()
            raise

    def get_batches(self, date_filter: str = None):
        if not self.conn:
            self.connect()
        if date_filter:
            query = """
                SELECT batch_id, merge_timestamp, box_count, valid_sn_count, total_sn_count, box_codes
                FROM merge_batches WHERE date(merge_timestamp) = ?
                ORDER BY merge_timestamp DESC
            """
            params = (date_filter,)
        else:
            query = """
                SELECT batch_id, merge_timestamp, box_count, valid_sn_count, total_sn_count, box_codes
                FROM merge_batches ORDER BY merge_timestamp DESC
            """
            params = ()
        cursor = self.conn.execute(query, params)
        results = []
        for row in cursor:
            batch_id, merge_timestamp, box_count, valid_sn_count, total_sn_count, box_codes_json = row
            try:
                box_codes = json.loads(box_codes_json)
                preview = (', '.join(box_codes[:3]) + '...') if len(box_codes) > 3 else ', '.join(box_codes)
            except:
                preview = box_codes_json[:50]
            results.append({
                'batch_id': batch_id,
                'merge_timestamp': merge_timestamp,
                'box_count': box_count,
                'valid_sn_count': valid_sn_count,
                'total_sn_count': total_sn_count,
                'box_codes_preview': preview
            })
        return results

    def get_all_versions(self):
        if not self.conn:
            self.connect()
        cursor = self.conn.execute("SELECT DISTINCT output_filename FROM merged_boxes")
        versions = set()
        for (output_filename,) in cursor:
            m = re.search(r'-v(\d+)', output_filename)
            if m:
                versions.add(m.group(1))
        return sorted(versions)

    def export_boxes_to_excel(self, output_path, date_filter: str = None,
                              version_filter: str = None,
                              upload_date_filter: str = None,
                              merge_date_filter: str = None) -> int:
        if not self.conn:
            self.connect()
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("BoxRecords")
        ws.append(["box_code", "original_filename", "shipment_date", "upload_time",
                   "merge_date", "merge_time", "output_filename",
                   "valid_rows", "total_rows", "sequence_no", "extracted_version"])
        count = 0
        query = ("SELECT box_code, original_filename, shipment_date, upload_time, merge_date, merge_time, "
                 "output_filename, valid_rows, total_rows, sequence_no, extracted_version "
                 "FROM merged_boxes WHERE 1=1")
        params = []
        if date_filter:
            query += " AND shipment_date = ?"; params.append(date_filter)
        if upload_date_filter:
            query += " AND upload_time LIKE ?"; params.append(f'{upload_date_filter}%')
        if merge_date_filter:
            query += " AND merge_date = ?"; params.append(merge_date_filter)
        if version_filter:
            query += " AND output_filename LIKE ?"; params.append(f'%-v{version_filter}%')
        query += " ORDER BY shipment_date, sequence_no"
        for row in self.conn.execute(query, params):
            ws.append(list(row)); count += 1
        wb.save(output_path)
        return count

    def export_sns_to_excel(self, output_path, date_filter: str = None,
                            version_filter: str = None,
                            upload_date_filter: str = None,
                            merge_date_filter: str = None) -> int:
        if not self.conn:
            self.connect()
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("SNRecords")
        ws.append(["SN", "box_code", "merge_date", "merge_time"])
        count = 0
        query = ("SELECT s.sn, s.box_code, s.merge_date, s.merge_time FROM merged_sns s "
                 "INNER JOIN merged_boxes b ON s.box_code = b.box_code WHERE 1=1")
        params = []
        if date_filter:
            query += " AND b.shipment_date = ?"; params.append(date_filter)
        if upload_date_filter:
            query += " AND b.upload_time LIKE ?"; params.append(f'{upload_date_filter}%')
        if merge_date_filter:
            query += " AND b.merge_date = ?"; params.append(merge_date_filter)
        if version_filter:
            query += " AND b.output_filename LIKE ?"; params.append(f'%-v{version_filter}%')
        query += " ORDER BY s.sn"
        for row in self.conn.execute(query, params):
            ws.append(list(row)); count += 1
        wb.save(output_path)
        return count

    def export_batches_to_excel(self, output_path, date_filter: str = None) -> int:
        if not self.conn:
            self.connect()
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("BatchRecords")
        ws.append(["batch_id", "merge_timestamp", "box_codes", "box_count",
                   "valid_sn_count", "total_sn_count",
                   "upload_start_time", "upload_end_time", "force_merge", "device_id"])
        count = 0
        if date_filter:
            query = ("SELECT batch_id, merge_timestamp, box_codes, box_count, valid_sn_count, total_sn_count, "
                     "upload_start_time, upload_end_time, force_merge, device_id "
                     "FROM merge_batches WHERE date(merge_timestamp) = ? ORDER BY merge_timestamp DESC")
            params = (date_filter,)
        else:
            query = ("SELECT batch_id, merge_timestamp, box_codes, box_count, valid_sn_count, total_sn_count, "
                     "upload_start_time, upload_end_time, force_merge, device_id "
                     "FROM merge_batches ORDER BY merge_timestamp DESC")
            params = ()
        for row in self.conn.execute(query, params):
            ws.append(list(row)); count += 1
        wb.save(output_path)
        return count

    def export_abnormal_to_excel(self, output_path, date_filter: str = None,
                                 upload_date_filter: str = None,
                                 merge_date_filter: str = None) -> int:
        if not self.conn:
            self.connect()
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("AbnormalData")
        ws.append(["sn", "box_code", "original_filename", "shipment_date",
                   "upload_time", "merge_date", "merge_time", "row_data"])
        count = 0
        query = ("SELECT sn, box_code, original_filename, shipment_date, upload_time, "
                 "merge_date, merge_time, row_data FROM abnormal_data WHERE 1=1")
        params = []
        if date_filter:
            query += " AND shipment_date = ?"; params.append(date_filter)
        if upload_date_filter:
            query += " AND upload_time LIKE ?"; params.append(f'{upload_date_filter}%')
        if merge_date_filter:
            query += " AND merge_date = ?"; params.append(merge_date_filter)
        query += " ORDER BY id"
        for row in self.conn.execute(query, params):
            ws.append(list(row)); count += 1
        wb.save(output_path)
        return count

    def export_skipped_to_excel(self, output_path, date_filter: str = None,
                                upload_date_filter: str = None,
                                merge_date_filter: str = None) -> int:
        if not self.conn:
            self.connect()
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("SkippedItems")
        ws.append(["item_type", "box_code", "sn", "source_file",
                   "reason_category", "reason_detail",
                   "shipment_date", "upload_time", "merge_date", "merge_time", "row_data"])
        count = 0
        query = ("SELECT item_type, box_code, sn, source_file, reason_category, reason_detail, "
                 "shipment_date, upload_time, merge_date, merge_time, row_data "
                 "FROM skipped_items WHERE 1=1")
        params = []
        if date_filter:
            query += " AND shipment_date = ?"; params.append(date_filter)
        if upload_date_filter:
            query += " AND upload_time LIKE ?"; params.append(f'{upload_date_filter}%')
        if merge_date_filter:
            query += " AND merge_date = ?"; params.append(merge_date_filter)
        query += " ORDER BY id"
        for row in self.conn.execute(query, params):
            ws.append(list(row)); count += 1
        wb.save(output_path)
        return count

    def get_database_info(self):
        if not self.conn:
            self.connect()
        boxes_count = self.conn.execute("SELECT COUNT(*) FROM merged_boxes").fetchone()[0]
        sns_count = self.conn.execute("SELECT COUNT(*) FROM merged_sns").fetchone()[0]
        mtime = os.path.getmtime(self.db_path)
        mtime_str = datetime.datetime.fromtimestamp(mtime).strftime('%Y-%m-%d %H:%M:%S')
        return boxes_count, sns_count, mtime_str

    def get_latest_boxes(self, limit=10):
        if not self.conn:
            self.connect()
        cursor = self.conn.execute(
            "SELECT box_code, original_filename, shipment_date, upload_time, "
            "merge_date, merge_time, output_filename, valid_rows, total_rows, sequence_no, "
            "extracted_version "
            "FROM merged_boxes ORDER BY id DESC LIMIT ?", (limit,))
        columns = ["box_code", "original_filename", "shipment_date", "upload_time",
                   "merge_date", "merge_time", "output_filename", "valid_rows", "total_rows",
                   "sequence_no", "extracted_version"]
        return [dict(zip(columns, row)) for row in cursor.fetchall()]