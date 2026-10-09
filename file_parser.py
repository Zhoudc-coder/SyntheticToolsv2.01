import re
import csv
import datetime
from pathlib import Path
from dataclasses import dataclass
from openpyxl import load_workbook


@dataclass
class ParsedBoxInfo:
    box_code: str
    package_date: str      # YYYY-MM-DD
    sequence_no: int       # 当日序号
    version: str           # 版本号（例如 '11'，不带 v）
    date_from_mtime: bool = False   # 是否使用文件修改时间作为日期


def extract_version(stem: str) -> str | None:
    """
    按优先级从文件名中提取版本号。
    优先级从高到低：
      1. (vXX)   英文括号 + 小写 v + 两位数字
      2. （vXX）  中文括号 + 小写 v + 两位数字
      3. (VXX)   英文括号 + 大写 V + 两位数字
      4. （VXX）  中文括号 + 大写 V + 两位数字
      5. vXX     无括号 + 小写 v + 两位数字
      6. VXX     无括号 + 大写 V + 两位数字（从文件名的第20个字符开始识别）
    兜底：任意位数数字
    """
    m = re.search(r'\(v(\d{2})\)', stem)
    if m:
        return m.group(1)
    m = re.search(r'（v(\d{2})）', stem)
    if m:
        return m.group(1)
    m = re.search(r'\(V(\d{2})\)', stem)
    if m:
        return m.group(1)
    m = re.search(r'（V(\d{2})）', stem)
    if m:
        return m.group(1)
    m = re.search(r'v(\d{2})', stem)
    if m:
        return m.group(1)
    if len(stem) > 19:
        m = re.search(r'V(\d{2})', stem[19:])
        if m:
            return m.group(1)
    m = re.search(r'[vV](\d+)', stem)
    if m:
        return m.group(1)
    return None


def extract_box_code_and_version_from_data_column(file_path: Path,
                                                  max_scan_rows: int = 20) -> tuple[str | None, str | None]:
    """
    从数据文件的"箱码"列中提取第一个非空箱码及其版本号。
    查找范围：倒数第 1 列到倒数第 5 列。
    匹配策略：
      1. 优先精确匹配标题为"箱码"的列；
      2. 若标题精确匹配失败（如 CSV 中文列名乱码），遍历倒数第 1~5 列，
         逐列尝试从数据中提取版本号，选择第一个能成功提取的列。
    只读取标题行和前 max_scan_rows 行数据。
    返回 (box_code, version)；未找到时返回 (None, None)。
    """
    def _process(title_row, data_iter):
        title_list = list(title_row) if title_row is not None else []
        title_len = len(title_list)
        if title_len == 0:
            return None, None

        # 收集倒数 1~5 列的索引
        candidate_indices = []
        for offset in range(1, 6):
            idx = title_len - offset
            if idx < 0:
                break
            candidate_indices.append(idx)

        if not candidate_indices:
            return None, None

        # 预读最多 max_scan_rows 行到内存，便于多列分别检查
        data_rows = []
        for i, row in enumerate(data_iter):
            if i >= max_scan_rows:
                break
            data_rows.append(row)

        # 第一步：精确匹配 "箱码" 列
        target_col_idx = None
        for idx in candidate_indices:
            if idx >= title_len:
                continue
            cell_val = title_list[idx]
            if cell_val is not None and str(cell_val).strip() == '箱码':
                target_col_idx = idx
                break

        # 第二步：确定要扫描的列顺序
        # 若标题精确匹配成功，仅扫描该列；否则扫描所有候选列
        search_indices = [target_col_idx] if target_col_idx is not None else candidate_indices

        for idx in search_indices:
            for row in data_rows:
                if row is None or len(row) <= idx:
                    continue
                cell_val = row[idx]
                if cell_val is None:
                    continue
                box_code = str(cell_val).strip()
                if not box_code:
                    continue
                v = extract_version(box_code)
                if v is not None:
                    return box_code, v

        return None, None

    suffix = file_path.suffix.lower()

    if suffix == '.xlsx':
        try:
            wb = load_workbook(filename=str(file_path), read_only=True, data_only=True)
            try:
                ws = wb.active
                try:
                    title_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
                except StopIteration:
                    return None, None
                rows = ws.iter_rows(min_row=2, values_only=True)
                return _process(title_row, rows)
            finally:
                wb.close()
        except Exception:
            return None, None

    elif suffix == '.csv':
        # CSV 编码尝试：utf-8-sig（Excel 保存带 BOM）、utf-8、gbk、gb18030
        encodings = ['utf-8-sig', 'utf-8', 'gbk', 'gb18030']
        for enc in encodings:
            try:
                with open(file_path, 'r', encoding=enc, newline='') as f:
                    reader = csv.reader(f)
                    try:
                        title_row = next(reader)
                    except StopIteration:
                        return None, None
                    return _process(title_row, reader)
            except UnicodeDecodeError:
                continue
        return None, None

    return None, None


def parse_filename(file_path: Path) -> ParsedBoxInfo | None:
    """
    解析文件名，提取箱码、打包日期、当日序号和版本号。
    如果文件名中找不到日期，则回退使用文件的修改时间（mtime）。
    仅在文件名中包含版本号时才返回 ParsedBoxInfo，否则返回 None。
    """
    stem = file_path.stem

    # 1. 提取日期（优先匹配 -YYMMDD- 格式）
    date_match = re.search(r'-(\d{6})-', stem)
    if not date_match:
        date_match = re.search(r'(\d{6})', stem)

    date_from_mtime = False
    package_date = None

    if date_match:
        date_str = date_match.group(1)
        yy, mm, dd = date_str[:2], date_str[2:4], date_str[4:6]
        try:
            year = 2000 + int(yy)
            month, day = int(mm), int(dd)
            package_date = datetime.date(year, month, day).isoformat()
        except ValueError:
            package_date = None

    # 2. 如果日期解析失败，使用文件修改时间
    if package_date is None:
        try:
            mtime = file_path.stat().st_mtime
            dt = datetime.datetime.fromtimestamp(mtime)
            package_date = dt.date().isoformat()
            date_from_mtime = True
        except Exception:
            return None

    # 3. 箱码使用完整文件名（不含扩展名）
    box_code = stem

    # 4. 提取序号（优先从日期后的部分提取，否则从整个文件名提取）
    if date_match:
        suffix = stem[date_match.end():]
    else:
        suffix = stem
    number_sequences = re.findall(r'\d+', suffix)
    if not number_sequences:
        return None

    sequence_no = None
    for num_str in reversed(number_sequences):
        if len(num_str) <= 4:
            try:
                sequence_no = int(num_str)
                break
            except ValueError:
                continue
    if sequence_no is None:
        try:
            sequence_no = int(number_sequences[0])
        except ValueError:
            return None

    # 5. 提取版本号（按优先级）
    version = extract_version(stem)
    if version is None:
        return None

    return ParsedBoxInfo(
        box_code=box_code,
        package_date=package_date,
        sequence_no=sequence_no,
        version=version,
        date_from_mtime=date_from_mtime
    )


def validate_sn(sn: str) -> bool:
    """
    校验SN格式：
    前半部分18位码，以CYV开头，以H3K结尾
    后半部分为八个以+分隔的字段，每个字段可为任意非空字符。
    """
    if not sn:
        return False
    sn = sn.strip()

    parts = sn.split('+')
    if len(parts) != 9:
        return False

    prefix = parts[0]
    if len(prefix) != 18:
        return False
    if not prefix.startswith('CYV') or not prefix.endswith('H3K'):
        return False

    for part in parts[1:]:
        if not part:
            return False

    return True